#!/usr/bin/env python3
"""Export PowerPoint decks to PDF by rendering each slide as an image-based PDF page."""

from pptx import Presentation
from pptx.util import Inches, Emu
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
import os, sys

BASE = "/Users/esam/Desktop/MatchingEngine/distribution"
DARK_NAVY = (26, 26, 46)
BLUE = (26, 35, 126)
WHITE = (255, 255, 255)
GRAY = (102, 102, 102)
LIGHT_GRAY = (153, 153, 153)
RED = (204, 0, 0)
AMBER = (214, 174, 0)


def hex_to_rgb(hex_str):
    """Convert hex color string to RGB tuple."""
    hex_str = hex_str.lstrip('#')
    return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))


def get_shape_text(shape):
    """Extract all text from a shape."""
    if not shape.has_text_frame:
        return ""
    lines = []
    for para in shape.text_frame.paragraphs:
        lines.append(para.text)
    return "\n".join(lines)


def render_slide_to_image(slide, width=1920, height=1440):
    """Render a slide to a PIL Image with text and embedded images."""
    img = Image.new('RGB', (width, height), WHITE)
    draw = ImageDraw.Draw(img)

    # Check for dark background slides (title/section dividers)
    is_dark = False
    for shape in slide.shapes:
        # Check if this slide has a dark background by looking at the background fill
        pass

    # Check slide background
    bg = slide.background
    if bg.fill.type is not None:
        try:
            rgb = bg.fill.fore_color.rgb
            r, g, b = int(str(rgb)[:2], 16), int(str(rgb)[2:4], 16), int(str(rgb)[4:6], 16)
            img = Image.new('RGB', (width, height), (r, g, b))
            draw = ImageDraw.Draw(img)
            is_dark = (r + g + b) < 200
        except:
            pass

    # Try to load fonts
    try:
        font_large = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 36)
        font_medium = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 22)
        font_small = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 18)
        font_tiny = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 14)
    except:
        font_large = ImageFont.load_default()
        font_medium = ImageFont.load_default()
        font_small = ImageFont.load_default()
        font_tiny = ImageFont.load_default()

    # Standard 10" x 7.5" slide in EMU
    slide_w = 9144000  # 10 inches
    slide_h = 6858000  # 7.5 inches
    scale_x = width / slide_w
    scale_y = height / slide_h

    for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):
        left = int((shape.left or 0) * scale_x)
        top = int((shape.top or 0) * scale_y)
        w = int((shape.width or 0) * scale_x)
        h = int((shape.height or 0) * scale_y)

        # Draw embedded pictures
        if shape.shape_type == 13:  # Picture
            try:
                from io import BytesIO
                image_stream = BytesIO(shape.image.blob)
                pic = Image.open(image_stream)
                pic = pic.resize((max(w, 1), max(h, 1)), Image.LANCZOS)
                img.paste(pic, (left, top))
            except Exception as e:
                # Draw placeholder
                draw.rectangle([left, top, left+w, top+h], outline=(200, 200, 200))
                draw.text((left+5, top+5), "[Image]", fill=GRAY, font=font_tiny)
            continue

        # Draw tables
        if shape.has_table:
            table = shape.table
            rows = len(table.rows)
            cols = len(table.columns)
            if rows > 0 and cols > 0:
                cell_w = w // cols
                cell_h = h // rows
                for r in range(rows):
                    for c in range(cols):
                        cell = table.cell(r, c)
                        cx = left + c * cell_w
                        cy = top + r * cell_h
                        # Header row
                        if r == 0:
                            draw.rectangle([cx, cy, cx+cell_w, cy+cell_h], fill=DARK_NAVY, outline=(200, 200, 200))
                            text_color = WHITE
                            font = font_tiny
                        else:
                            draw.rectangle([cx, cy, cx+cell_w, cy+cell_h], outline=(220, 220, 220))
                            text_color = GRAY
                            font = font_tiny
                        text = cell.text[:50]  # Truncate for display
                        draw.text((cx+4, cy+3), text, fill=text_color, font=font)
            continue

        # Draw text shapes
        if shape.has_text_frame:
            text = get_shape_text(shape)
            if not text.strip():
                continue

            # Determine font and color based on position and content
            text_color = WHITE if is_dark else DARK_NAVY
            font = font_medium

            # Title text (large, near top)
            if top < height * 0.15 and len(text) < 80:
                font = font_large
                text_color = WHITE if is_dark else DARK_NAVY
            # Planning assumption banners
            elif "Planning assumption" in text or "planning assumption" in text:
                font = font_tiny
                text_color = AMBER
            # Bullet content
            elif top > height * 0.15:
                font = font_small
                text_color = WHITE if is_dark else GRAY

            # Word wrap the text
            max_chars = max(w // 10, 20)
            lines = []
            for line in text.split('\n'):
                while len(line) > max_chars:
                    split_at = line[:max_chars].rfind(' ')
                    if split_at <= 0:
                        split_at = max_chars
                    lines.append(line[:split_at])
                    line = line[split_at:].lstrip()
                lines.append(line)

            y_pos = top
            for line in lines:
                if y_pos < height - 20:
                    draw.text((left, y_pos), line, fill=text_color, font=font)
                    y_pos += font.size + 6 if hasattr(font, 'size') else 20

    return img


def pptx_to_pdf(pptx_path, pdf_path):
    """Convert a PowerPoint file to PDF."""
    print(f"Converting: {os.path.basename(pptx_path)}")
    prs = Presentation(pptx_path)

    pdf = FPDF(orientation='L', unit='mm', format=(190.5, 254))  # 10" x 7.5" in mm

    for i, slide in enumerate(prs.slides):
        print(f"  Slide {i+1}/{len(prs.slides)}...", end="", flush=True)

        # Render slide to image
        slide_img = render_slide_to_image(slide, width=1920, height=1440)

        # Save temp image
        temp_path = f"/tmp/slide_{i}.png"
        slide_img.save(temp_path, "PNG")

        # Add to PDF
        pdf.add_page()
        pdf.image(temp_path, x=0, y=0, w=254, h=190.5)

        # Cleanup
        os.remove(temp_path)
        print(" done")

    pdf.output(pdf_path)
    print(f"  Created: {pdf_path} ({os.path.getsize(pdf_path) // 1024} KB)")


def main():
    decks = [
        ("Executive", "Matching-Engine-Executive-Summary-v1.0"),
        ("Architecture", "Matching-Engine-Architecture-Review-v1.0"),
    ]

    for deck_dir, deck_name in decks:
        pptx_path = f"{BASE}/{deck_dir}/{deck_name}.pptx"
        pdf_path = f"{BASE}/{deck_dir}/{deck_name}.pdf"
        if os.path.exists(pptx_path):
            pptx_to_pdf(pptx_path, pdf_path)
        else:
            print(f"MISSING: {pptx_path}")

    print("\nPDF export complete.")


if __name__ == "__main__":
    main()
